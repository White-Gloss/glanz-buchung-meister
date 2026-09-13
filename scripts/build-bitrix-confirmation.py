"""Build the native Bitrix confirmation template; fill only after a reserved approval."""
from pathlib import Path
import io
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

root = Path(__file__).resolve().parents[1]
logo = (root / 'docs/bitrix/templates/white-gloss-logo.png').read_bytes()
doc = Document()
# Remove inherited title rules from the runtime's default template.
for border in list(doc.styles.element.xpath('.//w:pBdr')):
    border.getparent().remove(border)
section = doc.sections[0]
# Match the existing White-Gloss A4 business stationery.
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin = section.bottom_margin = Inches(.65)
section.left_margin = section.right_margin = Inches(.85)
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'Arial', Pt(10)
normal.paragraph_format.space_after = Pt(7)
for name in ['Title', 'Heading 1', 'Heading 2']:
    doc.styles[name].font.name = 'Arial'
    doc.styles[name].font.color.rgb = RGBColor(0,0,0)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
# Preserve the supplied transparent logo; give its white lettering contrast.
shade = OxmlElement('w:shd')
shade.set(qn('w:fill'), '111111')
p._p.get_or_add_pPr().append(shade)
p.add_run().add_picture(io.BytesIO(logo), width=Inches(2.5))
doc.add_paragraph('Buchungsbestätigung', 'Title')
p = doc.add_paragraph('White Gloss Detailing · Lars Hägele\nArnistal 27 · 72160 Horb am Neckar')
p.runs[0].font.size = Pt(8)
doc.add_paragraph('Buchungsreferenz {WGReference}   ·   Ausgestellt am {WGIssuedOn}')
doc.add_paragraph('{WGCustomerName}\n{WGCustomerAddress}\n{WGCustomerEmail} · {WGCustomerPhone}')
doc.add_paragraph('Vielen Dank für deine Buchung. Nach unserer Prüfung bestätigen wir die folgenden Leistungen und den vereinbarten Termin.')
doc.add_paragraph('Fahrzeug', 'Heading 2')
doc.add_paragraph('{WGVehicle}\nKennzeichen {WGPlate}')
doc.add_paragraph('Bestätigte Leistungen', 'Heading 2')
doc.add_paragraph('{WGServiceLines}')
p = doc.add_paragraph('Endgültig vereinbarter Gesamtpreis: ')
p.add_run('{WGAgreedPrice}').bold = True
doc.add_paragraph('Vereinbarter Zeitraum', 'Heading 2')
doc.add_paragraph('Beginn: {WGStart}\nEnde: {WGEnd}\nVereinbarte Dauer: {WGDuration}\nLeistungsort: {WGLocation}')
doc.add_paragraph('Hinweise', 'Heading 2')
doc.add_paragraph('{WGNotes}')
p = doc.add_paragraph('Diese Buchungsbestätigung ist keine Rechnung. Die Rechnung wird nach der Durchführung der Dienstleistung separat erstellt.')
p.runs[0].font.size = Pt(9)
footer=section.footer.paragraphs[0]
footer.text='White Gloss Detailing · buchung@white-gloss.de · white-gloss.de'
footer.runs[0].font.size=Pt(8)
path=root/'docs/bitrix/templates/buchungsbestaetigung.docx'
doc.save(path)
print(path)
